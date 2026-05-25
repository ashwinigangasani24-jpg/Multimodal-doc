import os
import shutil
import fitz  # PyMuPDF
import pytesseract
import cv2
import numpy as np
import pandas as pd
from docx import Document
from io import BytesIO

# Configure Tesseract executable path for Windows, with environment override.
# If Tesseract is not installed or the path is wrong, a clear error is raised.
def _configure_tesseract_cmd():
    default_windows_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    configured_cmd = os.environ.get('TESSERACT_CMD', default_windows_cmd if os.name == 'nt' else 'tesseract')

    if os.path.isabs(configured_cmd):
        if os.path.exists(configured_cmd):
            pytesseract.pytesseract.tesseract_cmd = configured_cmd
            return
    elif shutil.which(configured_cmd):
        pytesseract.pytesseract.tesseract_cmd = configured_cmd
        return

    raise FileNotFoundError(
        f"Tesseract OCR executable not found. Please install Tesseract and set the full path to the executable via the environment variable TESSERACT_CMD. "
        f"Tried: {configured_cmd}. On Windows, install from https://github.com/UB-Mannheim/tesseract/wiki and set TESSERACT_CMD='C:\\Program Files\\Tesseract-OCR\\tesseract.exe'."
    )

if os.name == 'nt':
    _configure_tesseract_cmd()


def parse_pdf(file_path):
    """
    Parses a PDF file to extract text, tables, and images.
    """
    doc = fitz.open(file_path)
    text_content = []
    tables = []
    images = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        
        # Extract text
        text = page.get_text()
        if text.strip():
            text_content.append(text)
        
        # Extract tables (very basic extraction via PyMuPDF)
        tabs = page.find_tables()
        if tabs:
            for tab in tabs:
                df = pd.DataFrame(tab.extract())
                # Only append if we got a valid table
                if not df.empty:
                    tables.append(df)
                    
        # Extract images
        image_list = page.get_images(full=True)
        for img_index, img in enumerate(image_list):
            xref = img[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]
            images.append(image_bytes)

    # If text is extremely small, it might be a scanned PDF. Attempt OCR.
    full_text = "\n".join(text_content)
    if len(full_text.strip()) < 50 and images:
        full_text += "\n" + ocr_images_from_pdf(file_path)

    return {
        "text": full_text,
        "tables": tables,
        "images": images
    }

def ocr_images_from_pdf(file_path):
    """
    Helper function to run OCR on all pages of a PDF.
    Converts PDF pages to images using PyMuPDF and runs Tesseract.
    """
    doc = fitz.open(file_path)
    ocr_text = []
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
        if pix.n == 4:
            img = cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
        
        # Preprocessing for better OCR
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        # Apply thresholding
        _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY)
        
        text = pytesseract.image_to_string(thresh)
        ocr_text.append(text)
        
    return "\n".join(ocr_text)

def parse_docx(file_path):
    """
    Parses a DOCX file to extract text.
    """
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        text.append(para.text)
    
    # We can also extract basic tables
    tables = []
    for table in doc.tables:
        data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            data.append(row_data)
        if data:
            tables.append(pd.DataFrame(data[1:], columns=data[0]))
            
    return {
        "text": "\n".join(text),
        "tables": tables,
        "images": [] # DOCX image extraction is more complex, skipped for simplicity
    }

def parse_image(file_path):
    """
    Parses an image file using OCR to extract text.
    """
    img = cv2.imread(file_path)
    if img is None:
        return {"text": "", "tables": [], "images": []}
        
    # Preprocessing
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY)
    
    text = pytesseract.image_to_string(thresh)
    
    # Return the raw image bytes as well
    with open(file_path, "rb") as f:
        image_bytes = f.read()
        
    return {
        "text": text,
        "tables": [],
        "images": [image_bytes]
    }

def process_document(file_path, filename):
    """
    Main entry point for document parsing based on file extension.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    elif ext in ['.png', '.jpg', '.jpeg']:
        return parse_image(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
